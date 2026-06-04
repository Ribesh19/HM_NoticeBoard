#!/usr/bin/env python3
"""Build the Hockley Mint starter notice-board deck (board.pptx).

Recreates the eight default slides using the official brand palette, so the
manager never starts from a blank deck. 16:9, 1920x1080.

Backgrounds are built from the brand patterns (assets/pattern-01/02.png):
  - light slides: ivory field woven with pattern-02's green diamonds
  - cover/celebration slides: brand Dark Green textured with pattern-01
Text colours adapt to each background. The top-right corner is kept clear for
the live clock/weather overlay the kiosk draws on top.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

try:
    from PIL import Image, ImageChops
except ImportError:
    Image = None

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.normpath(os.path.join(HERE, "..", "assets"))

# ---- official brand palette ---------------------------------------------
BRIGHT = RGBColor(0x00, 0xA4, 0x78)   # Bright Green  (Pantone 3395 C)
DARKG  = RGBColor(0x00, 0x3D, 0x22)   # Dark Green    (Pantone 3435 C)
IVORY  = RGBColor(0xFD, 0xF3, 0xD6)   # Ivory         (Pantone 7499 C)
GREY_LN = RGBColor(0xC7, 0xC9, 0xBE)  # neutral rule (document mock only)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

HEADER = "Georgia"
BODY   = "Arial"

W = Inches(13.333)
H = Inches(7.5)
LM = Inches(0.95)   # left margin

# ---- per-background theming ---------------------------------------------
THEMES = {
    "light": dict(kicker=BRIGHT, title=DARKG, body=DARKG,
                  meta_label=BRIGHT, meta_value=DARKG,
                  card_fill=BRIGHT, card_text=IVORY, card_sub=IVORY,
                  chip_fill=DARKG, chip_text=IVORY, accent=BRIGHT),
    "green": dict(kicker=IVORY, title=IVORY, body=IVORY,
                  meta_label=BRIGHT, meta_value=IVORY,
                  card_fill=DARKG, card_text=IVORY, card_sub=IVORY,
                  chip_fill=IVORY, chip_text=DARKG, accent=BRIGHT),
}

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


# ---- background generation ----------------------------------------------
def _ivory_bg():
    """Ivory field (brand ivory) woven with pattern-02's green diamonds."""
    if Image is None:
        return None
    w, h = 1280, 720
    ivory = (0xFD, 0xF3, 0xD6)
    base = Image.new("RGB", (w, h), ivory)
    pat = os.path.join(ASSETS, "pattern-02.png")
    if os.path.exists(pat):
        p = Image.open(pat).convert("RGB").resize((w, h))
        diamonds = ImageChops.multiply(p, Image.new("RGB", (w, h), ivory))
        # keep it a soft tint, not a bold band, so dark-green text stays legible
        base = Image.blend(base, diamonds, 0.45)
    out = os.path.join(HERE, "_bg_light.png")
    base.save(out, optimize=True)
    return out


def _green_bg():
    """Brand Dark Green gradient textured faintly with pattern-01 diamonds."""
    if Image is None:
        return None
    w, h = 1280, 720
    dark, mid = (0x00, 0x3D, 0x22), (0x00, 0x5A, 0x38)
    grad = Image.new("RGB", (1, h))
    for y in range(h):
        t = y / (h - 1)
        grad.putpixel((0, y), tuple(int(dark[i] + (mid[i] - dark[i]) * t) for i in range(3)))
    base = grad.resize((w, h))
    pat = os.path.join(ASSETS, "pattern-01.png")
    if os.path.exists(pat):
        p = Image.open(pat).convert("RGB").resize((w, h))
        base = Image.blend(base, ImageChops.multiply(p, base), 0.5)
    out = os.path.join(HERE, "_bg_green.png")
    base.save(out, optimize=True)
    return out


BG_LIGHT = _ivory_bg()
BG_GREEN = _green_bg()


# ---- helpers -------------------------------------------------------------
def slide():
    return prs.slides.add_slide(BLANK)


def no_shadow(shape):
    sp = shape._element.spPr
    if sp.find(qn("a:effectLst")) is None:
        sp.append(sp.makeelement(qn("a:effectLst"), {}))


def bg(s, theme_name):
    img = BG_LIGHT if theme_name == "light" else BG_GREEN
    if img and os.path.exists(img):
        s.shapes.add_picture(img, 0, 0, width=W, height=H)
        return
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.line.fill.background(); no_shadow(r)
    r.fill.solid(); r.fill.fore_color.rgb = IVORY if theme_name == "light" else DARKG


def set_tracking(run, pts):
    run._r.get_or_add_rPr().set("spc", str(int(pts * 100)))


def text(s, l, t, w, h, runs, size, color, font=BODY, bold=False, italic=False,
         align=PP_ALIGN.LEFT, tracking=0.0, line_spacing=1.05, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [runs]
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for txt in runs:
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.name = font
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        if tracking:
            set_tracking(r, tracking)
    return box


def kicker(s, l, t, label, color):
    return text(s, l, t, Inches(9), Inches(0.5), label.upper(), 15, color,
                font=BODY, bold=True, tracking=4.0)


def card(s, l, t, w, h, fill=None, border=None, radius=0.08, line_w=1.2):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = radius
    no_shadow(sh)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border; sh.line.width = Pt(line_w)
    return sh


def circle(s, l, t, d, fill, glyph=None, glyph_color=IVORY, glyph_size=18, glyph_font=BODY):
    e = s.shapes.add_shape(MSO_SHAPE.OVAL, l, t, d, d)
    no_shadow(e)
    e.fill.solid(); e.fill.fore_color.rgb = fill
    e.line.fill.background()
    if glyph is not None:
        tf = e.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = glyph
        r.font.size = Pt(glyph_size); r.font.bold = True
        r.font.name = glyph_font; r.font.color.rgb = glyph_color
    return e


def meta(s, l, t, label, value, label_color, value_color, w=Inches(4)):
    text(s, l, t, w, Inches(0.4), label.upper(), 13, label_color, bold=True, tracking=3.0)
    text(s, l, t + Inches(0.38), w, Inches(0.6), value, 24, value_color,
         font=HEADER, bold=True)


# ========================================================================
# Slide 1 - Welcome (cover, green)
# ========================================================================
T = THEMES["green"]
s = slide()
bg(s, "green")
kicker(s, LM, Inches(1.15), "Today at Hockley Mint", T["kicker"])
text(s, LM, Inches(1.65), Inches(8.0), Inches(2.2),
     "Good morning, co-owners", 50, T["title"], font=HEADER, bold=True, line_spacing=1.0)
text(s, LM, Inches(3.55), Inches(6.7), Inches(1.8),
     "A quick look at the notices, events, birthdays, values and useful "
     "updates for the workshop today.", 22, T["body"], line_spacing=1.25)
meta(s, LM, Inches(5.55), "Workshop hours", "07:30 - 17:00", T["meta_label"], T["meta_value"])
logo = os.path.join(ASSETS, "logo-white.png")
if os.path.exists(logo):
    s.shapes.add_picture(logo, Inches(9.15), Inches(2.95), width=Inches(3.4))


# ========================================================================
# Slide 2 - Announcements (light)
# ========================================================================
T = THEMES["light"]
s = slide()
bg(s, "light")
kicker(s, LM, Inches(0.95), "Announcements", T["kicker"])
text(s, LM, Inches(1.45), Inches(7.4), Inches(1.6),
     "Annual co-owners' meeting", 46, T["title"], font=HEADER, bold=True, line_spacing=1.0)
text(s, LM, Inches(3.05), Inches(7.2), Inches(1.6),
     "All co-owners are invited to the main workshop. Agenda: FY results, "
     "Aurora collection, sustainability update and open Q&A.", 20, T["body"], line_spacing=1.25)
card(s, Inches(9.35), Inches(2.3), Inches(3.1), Inches(1.5), fill=T["card_fill"], radius=0.1)
text(s, Inches(9.65), Inches(2.55), Inches(2.6), Inches(0.4),
     "When", 13, T["card_text"], bold=True, tracking=3.0)
text(s, Inches(9.65), Inches(2.95), Inches(2.6), Inches(0.8),
     "Friday 16 May\n15:00", 24, T["card_text"], font=HEADER, bold=True, line_spacing=1.05)
items = ["Refreshments from 14:30", "RSVP on the kitchen sign-up sheet",
         "Questions welcome in advance"]
y = 4.85
for i, it in enumerate(items):
    text(s, LM, Inches(y), Inches(0.9), Inches(0.6),
         f"{i+1:02d}", 26, T["accent"], font=HEADER, bold=True)
    text(s, LM + Inches(0.85), Inches(y + 0.06), Inches(6.6), Inches(0.6),
         it, 20, T["body"], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.72


# ========================================================================
# Slide 3 - Birthday (celebration, green)
# ========================================================================
T = THEMES["green"]
s = slide()
bg(s, "green")
for (cx, cy, d, col) in [(3.8, 0.6, 0.16, BRIGHT), (5.0, 0.5, 0.10, IVORY),
                          (6.2, 0.78, 0.12, BRIGHT), (7.4, 0.55, 0.10, IVORY),
                          (8.4, 0.7, 0.09, BRIGHT)]:
    circle(s, Inches(cx), Inches(cy), Inches(d), col)
kicker(s, LM, Inches(1.5), "Birthday celebration", T["kicker"])
text(s, LM, Inches(2.0), Inches(7.4), Inches(1.8),
     "Happy Birthday, Aisha", 58, T["title"], font=HEADER, bold=True, line_spacing=1.0)
text(s, LM, Inches(4.2), Inches(6.7), Inches(1.7),
     "Wishing you a wonderful day from everyone at Hockley Mint. Thank you "
     "for the care and craft you bring to the workshop.", 22, T["body"], line_spacing=1.25)
meta(s, LM, Inches(6.0), "From", "Your Hockley Mint team", T["meta_label"], T["meta_value"], w=Inches(5.5))
ph = card(s, Inches(9.2), Inches(2.45), Inches(3.25), Inches(4.0), fill=None, border=IVORY, radius=0.06, line_w=1.5)
ph.line.dash_style = 2
text(s, Inches(9.2), Inches(4.2), Inches(3.25), Inches(0.6),
     "Add a photo", 16, IVORY, align=PP_ALIGN.CENTER, tracking=2.0)


# ========================================================================
# Slide 4 - Events (light)
# ========================================================================
T = THEMES["light"]
s = slide()
bg(s, "light")
kicker(s, LM, Inches(0.95), "Upcoming events", T["kicker"])
text(s, LM, Inches(1.45), Inches(8.5), Inches(1.2),
     "This fortnight at the workshop", 46, T["title"], font=HEADER, bold=True, line_spacing=1.0)
text(s, LM, Inches(3.05), Inches(8.0), Inches(0.8),
     "Key dates and useful reminders for everyone on site.", 20, T["body"])
events = [("08 May", "Co-owners coffee  -  Kitchen  -  09:00"),
          ("13 May", "Fire drill  -  whole building  -  11:30"),
          ("22 May", "Aurora press preview  -  Showroom  -  10:00")]
y = 3.75
for date, desc in events:
    card(s, LM, Inches(y), Inches(2.0), Inches(0.78), fill=T["chip_fill"], radius=0.18)
    text(s, LM, Inches(y), Inches(2.0), Inches(0.78), date, 22, T["chip_text"],
         font=HEADER, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, LM + Inches(2.35), Inches(y), Inches(8.4), Inches(0.78),
         desc, 22, T["body"], anchor=MSO_ANCHOR.MIDDLE)
    y += 1.02


# ========================================================================
# Slide 5 - Values (the real brand wallpaper, full-bleed)
# ========================================================================
s = slide()
values_img = os.path.join(ASSETS, "Values (Desktop Wallpaper).jpg")
if os.path.exists(values_img):
    s.shapes.add_picture(values_img, 0, 0, width=W, height=H)
else:
    T = THEMES["light"]
    bg(s, "light")
    kicker(s, LM, Inches(0.85), "Our values", T["kicker"])
    text(s, LM, Inches(1.35), Inches(9), Inches(1.1),
         "What we stand for", 46, T["title"], font=HEADER, bold=True)
    values = ["Kind to our environment", "Delighted customers",
              "Honest British craftsmanship", "Happy co-owners"]
    gx, gy = LM, Inches(2.75)
    cw, ch, gap = Inches(5.45), Inches(1.85), Inches(0.45)
    for i, v in enumerate(values):
        col, row = i % 2, i // 2
        l = gx + col * (cw + gap); t = gy + row * (ch + gap)
        card(s, l, t, cw, ch, fill=T["card_fill"], radius=0.08)
        circle(s, l + Inches(0.35), t + Inches(0.52), Inches(0.8), DARKG,
               glyph="◆", glyph_color=IVORY, glyph_size=22)
        text(s, l + Inches(1.45), t, cw - Inches(1.7), ch, v, 24, T["card_text"],
             font=HEADER, bold=True, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)


# ========================================================================
# Slide 6 - Health & Safety (light)
# ========================================================================
T = THEMES["light"]
s = slide()
bg(s, "light")
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), H)
no_shadow(bar); bar.fill.solid(); bar.fill.fore_color.rgb = T["accent"]; bar.line.fill.background()
kicker(s, LM, Inches(0.95), "Safety", T["kicker"])
text(s, LM, Inches(1.45), Inches(8.2), Inches(1.4),
     "Fire drill on Tuesday", 48, T["title"], font=HEADER, bold=True)
text(s, LM, Inches(2.95), Inches(7.4), Inches(1.4),
     "Scheduled drill for all departments. Please leave by your nearest route "
     "and meet at the east car park assembly point.", 20, T["body"], line_spacing=1.25)
card(s, Inches(9.5), Inches(2.2), Inches(3.0), Inches(1.6), fill=T["card_fill"], radius=0.1)
text(s, Inches(9.8), Inches(2.5), Inches(2.5), Inches(0.4),
     "Time", 13, T["card_text"], bold=True, tracking=3.0)
text(s, Inches(9.8), Inches(2.9), Inches(2.5), Inches(0.7),
     "Tuesday, 11:30", 22, T["card_text"], font=HEADER, bold=True, line_spacing=1.05)
items = ["Do not use lifts", "Take visitors with you",
         "Return only when cleared by Facilities"]
y = 4.7
for it in items:
    circle(s, LM, Inches(y + 0.05), Inches(0.32), T["accent"], glyph="✓",
           glyph_color=IVORY, glyph_size=14)
    text(s, LM + Inches(0.55), Inches(y), Inches(7.6), Inches(0.5),
         it, 20, T["body"], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.62


# ========================================================================
# Slide 7 - Kudos (people, light)
# ========================================================================
T = THEMES["light"]
s = slide()
bg(s, "light")
kicker(s, LM, Inches(0.9), "Shout-outs", T["kicker"])
text(s, LM, Inches(1.4), Inches(9), Inches(1.1),
     "Kudos from the workshop", 46, T["title"], font=HEADER, bold=True)
text(s, LM, Inches(2.6), Inches(9), Inches(0.7),
     "Small notes of appreciation from around the business.", 20, T["body"])
people = [("C", "Casting team", "Q1 wedding band run shipped a day early"),
          ("C", "CAD team", "Bespoke commission turned around in 48 hours"),
          ("D", "Despatch", "Every next-day parcel out before cut-off")]
cw, gap = Inches(3.7), Inches(0.42)
x = LM
for initial, who, what in people:
    card(s, x, Inches(3.5), cw, Inches(2.95), fill=T["card_fill"], radius=0.08)
    circle(s, x + Inches(0.4), Inches(3.9), Inches(0.85), DARKG,
           glyph=initial, glyph_color=IVORY, glyph_size=26, glyph_font=HEADER)
    text(s, x + Inches(0.4), Inches(4.95), cw - Inches(0.8), Inches(0.45),
         who, 20, T["card_text"], font=HEADER, bold=True)
    text(s, x + Inches(0.4), Inches(5.45), cw - Inches(0.8), Inches(0.9),
         what, 16, T["card_sub"], line_spacing=1.2)
    x += cw + gap


# ========================================================================
# Slide 8 - Document (light)
# ========================================================================
T = THEMES["light"]
s = slide()
bg(s, "light")
kicker(s, LM, Inches(1.2), "Useful document", T["kicker"])
text(s, LM, Inches(1.7), Inches(7.4), Inches(1.3),
     "Brand Guidelines V2", 48, T["title"], font=HEADER, bold=True)
text(s, LM, Inches(3.15), Inches(6.7), Inches(1.5),
     "The latest brand guidelines are available for anyone preparing internal "
     "or customer-facing materials.", 22, T["body"], line_spacing=1.25)
meta(s, LM, Inches(5.2), "PDF", "P4707 Hockley Mint Brand Guidelines V2",
     T["meta_label"], T["meta_value"], w=Inches(7))
sheet = card(s, Inches(9.35), Inches(2.35), Inches(3.0), Inches(3.9),
             fill=WHITE, border=DARKG, radius=0.03, line_w=1.0)
corner = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(11.55), Inches(2.35), Inches(0.8), Inches(0.55))
no_shadow(corner); corner.fill.solid(); corner.fill.fore_color.rgb = BRIGHT; corner.line.fill.background()
text(s, Inches(11.55), Inches(2.35), Inches(0.8), Inches(0.55), "PDF", 13, IVORY,
     bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
ly = 3.25
for w_in in [2.2, 2.4, 1.8, 2.4, 2.1, 1.6]:
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.7), Inches(ly), Inches(w_in), Inches(0.12))
    no_shadow(ln); ln.fill.solid(); ln.fill.fore_color.rgb = GREY_LN; ln.line.fill.background()
    ly += 0.42

out = os.path.join(HERE, "board.pptx")
prs.save(out)
print("wrote", out, "with", len(prs.slides._sldIdLst), "slides")
